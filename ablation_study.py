#!/usr/bin/env python3
"""
HikmaProf Ablation Study
Automatically generates presentations, evaluates them, and outputs a final table.

Usage:
    1. Set your PDF paths in PDF_DOCUMENTS below
    2. Make sure app.py is running on APP_URL
    3. Run: python ablation_study.py
"""

import os
import time
import requests
import pandas as pd
import textstat
import fitz                          # PyMuPDF
from pathlib import Path
from pptx import Presentation
from rouge_score import rouge_scorer as rouge_lib

# ── Configuration ──────────────────────────────────────────────────────────────

APP_URL = "http://localhost:9000"    # HikmaProf server

import glob
PDF_DOCUMENTS = sorted(glob.glob("doc/*.pdf"))

VARIANTS = [
    {"name": "T1 - Modern (No AI Images)",    "template": "modern",    "color": "blue", "num_images": 0},
    {"name": "T2 - Minimal (No AI Images)",   "template": "minimal",   "color": "blue", "num_images": 0},
    {"name": "T3 - Corporate (No AI Images)", "template": "corporate", "color": "blue", "num_images": 0},
    {"name": "T4 - Modern + 4 AI Images",     "template": "modern",    "color": "blue", "num_images": 4},
    {"name": "T5 - Modern + 5 AI Images",     "template": "modern",    "color": "blue", "num_images": 5},
]

OUTPUT_DIR = "ablation_results"
os.makedirs(OUTPUT_DIR, exist_ok=True)


# ── Text Extraction ────────────────────────────────────────────────────────────

def extract_pdf_text(pdf_path):
    doc = fitz.open(pdf_path)
    return " ".join(page.get_text() for page in doc)

def extract_pptx_text(pptx_path):
    prs = Presentation(pptx_path)
    texts = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                texts.append(shape.text_frame.text)
    return " ".join(texts)


# ── Metrics ────────────────────────────────────────────────────────────────────

def compute_rouge_l(source_text, slide_text):
    scorer = rouge_lib.RougeScorer(['rougeL'])
    score  = scorer.score(source_text, slide_text)['rougeL'].fmeasure
    return round(score, 3)

def compute_flesch(text):
    return round(textstat.flesch_reading_ease(text), 1)


# ── HikmaProf API Calls ────────────────────────────────────────────────────────

def submit_job(pdf_path, variant):
    with open(pdf_path, 'rb') as f:
        files = {'file': (Path(pdf_path).name, f, 'application/pdf')}
        data  = {
            'template_style':  variant['template'],
            'color_scheme':    variant['color'],
            'num_images':      variant['num_images'],
            'skip_validation': 'false',
            'preview_slides':  'false',
            'num_slides':      12,
            'language':        'en',
            'avatar_choice':   '',
        }
        start = time.time()
        resp  = requests.post(f"{APP_URL}/api/generate", files=files, data=data, timeout=60)
    resp.raise_for_status()
    return resp.json()['job_id'], start

def wait_for_job(job_id, start_time, timeout=900):
    for _ in range(timeout // 3):
        time.sleep(3)
        job    = requests.get(f"{APP_URL}/api/job/{job_id}", timeout=10).json()
        status = job.get('status', '')
        print(f"    [{status}] {job.get('progress', 0)}%", end='\r')

        if status == 'completed':
            gen_time = round(time.time() - start_time, 1)
            pptx_file = job.get('output_file', '')
            return gen_time, pptx_file

        if status in ('failed', 'error'):
            raise Exception(job.get('error', 'Generation failed'))

    raise Exception("Timed out")

def download_pptx(job_id, pptx_file, save_path):
    resp = requests.get(f"{APP_URL}/api/download/{job_id}/{pptx_file}", timeout=60)
    resp.raise_for_status()
    with open(save_path, 'wb') as f:
        f.write(resp.content)


# ── Main Runner ────────────────────────────────────────────────────────────────

def run_ablation():
    all_results = []

    for doc_idx, pdf_path in enumerate(PDF_DOCUMENTS):
        if not os.path.exists(pdf_path):
            print(f"\n[SKIP] {pdf_path} — file not found")
            continue

        doc_name    = Path(pdf_path).name
        source_text = extract_pdf_text(pdf_path)

        print(f"\n{'='*65}")
        print(f"  Document {doc_idx+1}/{len(PDF_DOCUMENTS)}: {doc_name}")
        print(f"{'='*65}")

        for variant in VARIANTS:
            print(f"\n  Running: {variant['name']}")
            try:
                job_id, start = submit_job(pdf_path, variant)
                gen_time, pptx_file = wait_for_job(job_id, start)

                safe_name = variant['name'].replace(' ', '_').replace('/', '-')
                pptx_path = os.path.join(OUTPUT_DIR, f"doc{doc_idx+1}_{safe_name}.pptx")
                download_pptx(job_id, pptx_file, pptx_path)

                slide_text = extract_pptx_text(pptx_path)
                rl  = compute_rouge_l(source_text, slide_text)
                fk  = compute_flesch(slide_text)

                print(f"  OK Time: {gen_time}s  |  ROUGE-L: {rl}  |  Flesch: {fk}    ")
                all_results.append({
                    "Document":      doc_name,
                    "Variant":       variant['name'],
                    "Gen Time (s)":  gen_time,
                    "ROUGE-L":       rl,
                    "Flesch Score":  fk,
                })

            except Exception as e:
                print(f"  FAIL  Error: {e}    ")
                all_results.append({
                    "Document":     doc_name,
                    "Variant":      variant['name'],
                    "Gen Time (s)": None,
                    "ROUGE-L":      None,
                    "Flesch Score": None,
                })

    # ── Build Results Table ────────────────────────────────────────────────────
    df = pd.DataFrame(all_results)
    df.to_csv(os.path.join(OUTPUT_DIR, "ablation_raw.csv"), index=False)

    numeric = ["Gen Time (s)", "ROUGE-L", "Flesch Score"]
    summary = (df.groupby("Variant")[numeric]
                 .mean()
                 .round(3)
                 .reset_index())
    summary.columns = ["Variant", "Avg Gen Time (s)", "Avg ROUGE-L", "Avg Flesch"]
    summary.to_csv(os.path.join(OUTPUT_DIR, "ablation_summary.csv"), index=False)

    print("\n\n" + "="*75)
    print("  ABLATION STUDY — FINAL RESULTS (averaged across all documents)")
    print("="*75)
    print(summary.to_string(index=False))
    print(f"\nRaw results  → {OUTPUT_DIR}/ablation_raw.csv")
    print(f"Summary table → {OUTPUT_DIR}/ablation_summary.csv")

    return summary


if __name__ == "__main__":
    run_ablation()
